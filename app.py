import io
import os
import re
import json
from pathlib import Path

import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pypdf import PdfReader

st.set_page_config(page_title="AI Presentation Generator", page_icon="🎯", layout="wide")

# ---------- Professional theme ----------
NAVY = RGBColor(22, 40, 72)
BLUE = RGBColor(44, 105, 176)
TEAL = RGBColor(32, 158, 151)
PALE = RGBColor(239, 245, 252)
DARK = RGBColor(45, 55, 72)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(105, 119, 138)

st.markdown("""
<style>
    .stApp { background: #f5f7fb; }
    .hero { padding: 1.3rem 1.5rem; border-radius: 18px; color: white;
            background: linear-gradient(120deg, #162848, #2c69b0); margin-bottom: 1rem; }
    .hero h1 { color: white; margin: 0; }
    .hero p { color: #e4eefb; margin-bottom: 0; }
    div[data-testid="stMetric"] { background: white; padding: 12px; border-radius: 12px; }
</style>
<div class="hero">
  <h1>🎯 AI Presentation Generator</h1>
  <p>Turn a topic, your notes, or an uploaded document into a polished, editable PowerPoint.</p>
</div>
""", unsafe_allow_html=True)


def read_uploaded_file(uploaded):
    if uploaded is None:
        return ""
    name = uploaded.name.lower()
    try:
        if name.endswith(".txt") or name.endswith(".md"):
            return uploaded.getvalue().decode("utf-8", errors="ignore")
        if name.endswith(".pdf"):
            reader = PdfReader(uploaded)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if name.endswith(".docx"):
            from docx import Document
            doc = Document(uploaded)
            return "\n".join(p.text for p in doc.paragraphs)
    except Exception as exc:
        st.warning(f"Could not read this file: {exc}")
    return ""


def clean_lines(text):
    lines = []
    for line in text.splitlines():
        line = re.sub(r"^\s*(?:[-*•]+|\d+[.)])\s*", "", line).strip()
        if line:
            lines.append(line)
    return lines


def make_outline(topic, source_text, slide_count, audience, tone):
    """No-key fallback: creates a useful structured draft from the user's topic and source."""
    source_lines = clean_lines(source_text)
    sentences = []
    if source_text.strip():
        chunks = re.split(r"(?<=[.!?])\s+|\n+", source_text.strip())
        sentences = [s.strip(" •-*#\t") for s in chunks if len(s.strip()) > 18]

    title = topic.strip() or (source_lines[0] if source_lines else "My Presentation")
    title = title[:100]
    slides = [{"title": title, "bullets": [
        f"An overview of {title}",
        f"Designed for {audience.lower()}",
        f"Purpose: explain the key ideas and practical takeaways"
    ], "kind": "title"}]

    if slide_count <= 2:
        slides.append({"title": "Key takeaways", "bullets": [
            "Understand the main idea",
            "Connect the topic to practical examples",
            "Identify one action to take next"
        ], "kind": "content"})
        return slides[:slide_count]

    # Select useful source sentences first; if there is no source, use a topic-based scaffold.
    usable = sentences[:]
    if usable:
        # Spread source statements across requested slides.
        content_count = slide_count - 2
        per_slide = max(1, (len(usable) + content_count - 1) // content_count)
        for i in range(content_count):
            chosen = usable[i * per_slide:(i + 1) * per_slide]
            if not chosen:
                break
            slides.append({
                "title": f"Key idea {i + 1}",
                "bullets": [s[:220] for s in chosen[:4]],
                "kind": "content"
            })
    else:
        sections = [
            ("Why it matters", [
                f"How {title} affects people, organizations, or everyday decisions",
                "The problem or opportunity this topic helps us understand",
                "Why the topic is relevant today"
            ]),
            ("Core concepts", [
                f"The main principles behind {title}",
                "Important terms and ideas to know",
                "How the concepts connect to one another"
            ]),
            ("How it works", [
                "Break the process into clear, manageable steps",
                "Use examples to explain each step",
                "Check assumptions and results"
            ]),
            ("Examples and applications", [
                f"Practical ways to apply {title}",
                "A realistic example or short case study",
                "What a successful result looks like"
            ]),
            ("Benefits and opportunities", [
                "Potential improvements in quality, speed, or understanding",
                "Opportunities for learning and innovation",
                "Ways to measure whether it is helping"
            ]),
            ("Challenges to consider", [
                "Check accuracy and avoid unsupported claims",
                "Consider privacy, fairness, and accessibility",
                "Adapt the approach to the audience and context"
            ]),
            ("Recommendations", [
                "Start with a clear goal",
                "Test the approach on a small scale",
                "Use feedback to improve the result"
            ]),
        ]
        for i in range(slide_count - 2):
            section = sections[i % len(sections)]
            slides.append({"title": section[0], "bullets": section[1], "kind": "content"})

    slides.append({"title": "Conclusion", "bullets": [
        f"Remember the central message: {title}",
        "Focus on the most relevant ideas and evidence",
        "Questions and discussion"
    ], "kind": "closing"})
    slides = slides[:slide_count]
    if len(slides) < slide_count:
        slides.append({"title": "Questions & Discussion", "bullets": [
            "What is the most important insight?",
            "What would you like to explore further?"
        ], "kind": "closing"})
    return slides


def generate_with_gemini(api_key, topic, source_text, slide_count, audience, tone):
    """Optional free-tier Gemini API mode. The app still works without an API key."""
    import requests
    prompt = f"""
Create a presentation outline as valid JSON only. Do not use markdown fences.
Return this exact shape: {{"slides":[{{"title":"...", "bullets":["...", "..."]}}]}}
Create exactly {slide_count} slides about: {topic}
Audience: {audience}. Tone: {tone}.
Use concise, accurate, presentation-friendly bullets (3-5 bullets per slide, each under 22 words).
Use this source material when provided; do not invent specific statistics or citations:
{source_text[:12000]}
Make the first slide a title/overview slide and the last slide a conclusion or Q&A slide.
"""
    response = requests.post(
        "https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent",
        params={"key": api_key},
        json={"contents": [{"parts": [{"text": prompt}]}],
              "generationConfig": {"temperature": 0.4, "responseMimeType": "application/json"}},
        timeout=60,
    )
    response.raise_for_status()
    data = response.json()
    text = data["candidates"][0]["content"]["parts"][0]["text"]
    parsed = json.loads(text)
    slides = parsed.get("slides", [])
    if not slides:
        raise ValueError("The AI response did not contain slides.")
    clean = []
    for slide in slides[:slide_count]:
        bullets = slide.get("bullets", [])
        clean.append({"title": str(slide.get("title", "Slide"))[:100],
                      "bullets": [str(x)[:240] for x in bullets[:5]],
                      "kind": "content"})
    return clean


def add_text(slide, text, x, y, w, h, size, color, bold=False, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def build_powerpoint(slides, topic, presenter, theme_name):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for index, item in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        is_title = index == 0
        is_last = index == len(slides) - 1

        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = NAVY if is_title else WHITE

        # Decorative accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = TEAL if theme_name == "Teal" else BLUE
        bar.line.fill.background()

        if is_title:
            add_text(slide, "PRESENTATION", 0.85, 1.1, 4.0, 0.4, 14, TEAL, True)
            add_text(slide, item["title"], 0.85, 2.0, 11.3, 1.7, 34, WHITE, True)
            add_text(slide, "A clear, practical overview", 0.9, 4.0, 9.5, 0.6, 20, RGBColor(215, 229, 246))
            if presenter.strip():
                add_text(slide, presenter.strip(), 0.9, 6.2, 8.0, 0.4, 14, RGBColor(215, 229, 246))
            add_text(slide, f"{index + 1:02d}", 12.0, 6.8, 0.5, 0.3, 11, RGBColor(215, 229, 246))
            continue

        # Top title and divider
        add_text(slide, item["title"], 0.8, 0.55, 11.7, 0.8, 27, NAVY, True)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.82), Inches(1.45), Inches(1.15), Inches(0.06))
        line.fill.solid()
        line.fill.fore_color.rgb = TEAL if theme_name == "Teal" else BLUE
        line.line.fill.background()

        bullets = item.get("bullets", []) or ["Add your key point here."]
        top = 1.85
        usable_height = 4.65
        each_h = min(1.08, usable_height / max(len(bullets), 1))
        for b_index, bullet in enumerate(bullets[:5]):
            y = top + b_index * each_h
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.92), Inches(y + 0.12), Inches(0.14), Inches(0.14))
            dot.fill.solid()
            dot.fill.fore_color.rgb = TEAL if theme_name == "Teal" else BLUE
            dot.line.fill.background()
            add_text(slide, str(bullet), 1.25, y, 10.9, each_h - 0.05, 19, DARK)
        # Footer
        footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.85), Inches(11.7), Inches(0.012))
        footer.fill.solid()
        footer.fill.fore_color.rgb = PALE
        footer.line.fill.background()
        add_text(slide, topic[:70] if topic else "Presentation", 0.85, 6.95, 9.8, 0.25, 9, MUTED)
        add_text(slide, f"{index + 1:02d}", 11.9, 6.92, 0.5, 0.3, 10, MUTED)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output.getvalue()


def outline_to_text(slides):
    parts = []
    for i, slide in enumerate(slides, start=1):
        parts.append(f"SLIDE {i}: {slide['title']}")
        for bullet in slide.get("bullets", []):
            parts.append(f"  • {bullet}")
        parts.append("")
    return "\n".join(parts)


with st.sidebar:
    st.header("⚙️ Presentation settings")
    topic = st.text_input("Presentation topic / title", placeholder="e.g., Generative AI in Education")
    slide_count = st.slider("Number of slides", min_value=3, max_value=12, value=7)
    audience = st.selectbox("Audience", ["University students", "Teachers", "Business professionals", "General audience", "Clients"])
    tone = st.selectbox("Writing style", ["Professional", "Simple and beginner-friendly", "Persuasive", "Academic"])
    presenter = st.text_input("Presenter name (optional)")
    theme_name = st.selectbox("Color theme", ["Blue", "Teal"])

    st.divider()
    st.subheader("Optional AI enhancement")
    st.caption("The app works without a key using its built-in no-key generator. For richer AI writing, you can optionally use a Gemini API key from Google's free tier.")
    api_key = st.text_input("Gemini API key (optional)", type="password", help="Leave blank to use no-key mode. Never share your key.")
    use_ai = st.checkbox("Use Gemini AI if a key is entered", value=False)

left, right = st.columns([1.05, 0.95], gap="large")
with left:
    st.subheader("1. Add your content")
    source_mode = st.radio("Choose your input", ["Topic only", "Paste notes", "Upload a document", "Topic + notes/document"], horizontal=False)
    notes = st.text_area("Paste notes or source text", height=190, placeholder="Paste your class notes, research, outline, or key facts here...")
    uploaded = st.file_uploader("Upload a source file (PDF, TXT, MD, DOCX)", type=["pdf", "txt", "md", "docx"])
    extracted = read_uploaded_file(uploaded) if uploaded else ""
    if uploaded and extracted:
        st.success(f"Read text from {uploaded.name} ({len(extracted):,} characters).")
    elif uploaded:
        st.info("No readable text was extracted. Scanned PDFs may need OCR first.")

with right:
    st.subheader("2. Generate your presentation")
    st.write("Your output is an editable `.pptx` file. You can change the text, colors, layout, and slides in PowerPoint or compatible editors.")
    st.markdown("- **Topic only:** creates a structured starter deck.")
    st.markdown("- **Paste notes / upload:** uses your material to build slides.")
    st.markdown("- **No paid API required:** built-in mode works without a key.")
    generate = st.button("✨ Generate presentation", type="primary", use_container_width=True)

if generate:
    if source_mode in ["Topic only", "Topic + notes/document"] and not topic.strip() and not (notes.strip() or extracted.strip()):
        st.error("Please enter a topic or add some source content.")
        st.stop()
    combined_source = ""
    if source_mode == "Paste notes":
        combined_source = notes
    elif source_mode == "Upload a document":
        combined_source = extracted
    elif source_mode == "Topic + notes/document":
        combined_source = "\n".join(x for x in [notes, extracted] if x.strip())
    effective_topic = topic.strip()
    if not effective_topic and combined_source:
        first_line = clean_lines(combined_source)
        effective_topic = first_line[0][:100] if first_line else "My Presentation"

    with st.spinner("Building your slides..."):
        slides = None
        if use_ai and api_key.strip():
            try:
                slides = generate_with_gemini(api_key.strip(), effective_topic, combined_source, slide_count, audience, tone)
                st.success("Generated using Gemini AI.")
            except Exception as exc:
                st.warning(f"Gemini generation failed, so the app used the built-in generator instead. Details: {exc}")
        if slides is None:
            slides = make_outline(effective_topic, combined_source, slide_count, audience, tone)
            st.info("Generated with the built-in no-key mode. Review the draft before presenting.")
        # Normalize to exact requested slide count.
        slides = slides[:slide_count]
        while len(slides) < slide_count:
            slides.append({"title": "Additional points", "bullets": ["Add a supporting detail", "Include a relevant example"], "kind": "content"})
        pptx_bytes = build_powerpoint(slides, effective_topic, presenter, theme_name)
        outline_text = outline_to_text(slides)

    st.session_state["generated_slides"] = slides
    st.session_state["generated_pptx"] = pptx_bytes
    st.session_state["generated_outline"] = outline_text
    st.session_state["generated_topic"] = effective_topic

if "generated_slides" in st.session_state:
    st.divider()
    st.subheader("3. Review and download")
    slides = st.session_state["generated_slides"]
    m1, m2, m3 = st.columns(3)
    m1.metric("Slides", len(slides))
    m2.metric("Editable format", "PPTX")
    m3.metric("AI setup", "Optional")
    st.caption("This is a draft. Check facts, add citations where needed, and personalize the slides before presenting.")
    with st.expander("Preview slide outline", expanded=True):
        for i, slide in enumerate(slides, start=1):
            st.markdown(f"**Slide {i}: {slide['title']}**")
            for bullet in slide.get("bullets", []):
                st.markdown(f"- {bullet}")
    c1, c2 = st.columns(2)
    with c1:
        st.download_button(
            "⬇️ Download editable PowerPoint",
            data=st.session_state["generated_pptx"],
            file_name="ai_presentation.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            use_container_width=True,
        )
    with c2:
        st.download_button(
            "⬇️ Download outline (.txt)",
            data=st.session_state["generated_outline"].encode("utf-8"),
            file_name="presentation_outline.txt",
            mime="text/plain",
            use_container_width=True,
        )
else:
    st.info("Start by entering a topic in the left settings panel, then add notes or upload a document if you have one.")
